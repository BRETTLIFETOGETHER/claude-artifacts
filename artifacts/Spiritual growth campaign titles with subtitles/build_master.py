import json, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from new_collections import NEW_COLLECTIONS
from collections import defaultdict, Counter

OUT = "/mnt/user-data/outputs/LifeTogether-Grand-Master-Library.xlsx"

CODES = {
    "Spiritual Growth Library": "SGL",
    "Phase 1 Launch Library": "PLL",
    "Faith at Work & Affinity Channels": "FAW",
    "Themes & Formation": "TAF",
    "Top 100 Flagship": "FLG",
    "Preaching Forecast": "FOR",
    "Capital Campaign": "CAP",
    "Discipleship": "DSC",
    "Fellowship": "FEL",
    "Missions & Evangelism": "MIS",
    "Worship": "WOR",
    "Ministry & Service": "SRV",
}
ORDER = list(CODES.keys())

records = []

wb0 = openpyxl.load_workbook("/mnt/user-data/outputs/LifeTogether-Master-Title-Directory.xlsx", data_only=True)
ws0 = wb0["Master Directory"]
hdr = [c.value for c in ws0[1]]
idx = {h: i for i, h in enumerate(hdr)}
for row in ws0.iter_rows(min_row=2, values_only=True):
    coll = row[idx["Collection"]]
    if not coll:
        continue
    cat = row[idx["Category"]] or ""
    grp = row[idx["Group"]] or ""
    records.append({
        "collection": coll,
        "category": str(cat if cat else grp),
        "title": str(row[idx["Title"]] or "").strip(),
        "subtitle": str(row[idx["Subtitle"]] or "").strip(),
    })

html_rows = json.load(open("/home/claude/html_rows.json", encoding="utf-8"))
for r in html_rows:
    records.append({"collection": r["collection"], "category": r.get("category", ""),
                    "title": r["title"].strip(), "subtitle": r.get("subtitle", "").strip()})

for coll, items in NEW_COLLECTIONS.items():
    for title, sub in items:
        records.append({"collection": coll, "category": "", "title": title.strip(), "subtitle": sub.strip()})

def coll_rank(c):
    return ORDER.index(c) if c in ORDER else len(ORDER)
records = [r for _, r in sorted(enumerate(records), key=lambda kv: (coll_rank(kv[1]["collection"]), kv[0]))]

seq = {}
for r in records:
    code = CODES.get(r["collection"], "GEN")
    seq[code] = seq.get(code, 0) + 1
    r["id"] = f"{code}-{seq[code]:04d}"

def norm(t):
    t = t.lower().strip().replace("&", " and ")
    t = re.sub(r"[^a-z0-9 ]", "", t)
    return re.sub(r"\s+", " ", t).strip()

groups = defaultdict(list)
for i, r in enumerate(records):
    r["key"] = norm(r["title"])
    groups[r["key"]].append(i)

for key, idxs in groups.items():
    canon = idxs[0]
    colls = {records[j]["collection"] for j in idxs}
    for j in idxs:
        records[j]["occurrences"] = len(idxs)
        records[j]["canonical"] = "Yes" if j == canon else "No"
        records[j]["canonical_id"] = records[canon]["id"]
        records[j]["in_collections"] = len(colls)

total = len(records)
unique = len(groups)
dup_instances = sum(1 for r in records if r["canonical"] == "No")
recurring = sum(1 for k, v in groups.items() if len(v) > 1)
cross = sum(1 for k, v in groups.items() if len({records[j]["collection"] for j in v}) > 1)
top_dups = sorted(
    ({"title": records[v[0]]["title"], "occ": len(v),
      "colls": len({records[j]["collection"] for j in v})}
     for k, v in groups.items() if len(v) > 1),
    key=lambda d: (-d["occ"], d["title"]))[:25]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Master Directory"
cols = ["ID", "Collection", "Category", "Title", "Subtitle",
        "Occurrences", "Canonical?", "Canonical ID", "In # Collections", "Norm Key"]
ws.append(cols)
for r in records:
    ws.append([r["id"], r["collection"], r["category"], r["title"], r["subtitle"],
               r["occurrences"], r["canonical"], r["canonical_id"], r["in_collections"], r["key"]])

ARIAL = "Arial"
hf = PatternFill("solid", fgColor="2F5B3E")
for c in ws[1]:
    c.fill = hf; c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=11)
    c.alignment = Alignment(vertical="center")
ws.freeze_panes = "A2"
ndata = total + 1
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ndata}"
for row in ws.iter_rows(min_row=2, max_row=ndata):
    for c in row:
        c.font = Font(name=ARIAL, size=10)
for i, w in enumerate([12, 30, 26, 30, 46, 12, 11, 12, 15, 26], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ix = wb.create_sheet("Index", 0)
D = "'Master Directory'!"
IDR, COLLR, CANR, OCCR = f"{D}$A$2:$A${ndata}", f"{D}$B$2:$B${ndata}", f"{D}$G$2:$G${ndata}", f"{D}$F$2:$F${ndata}"
tf = Font(name=ARIAL, bold=True, size=16, color="2F5B3E")
h2 = Font(name=ARIAL, bold=True, size=11, color="1A2A20")
lbl = Font(name=ARIAL, size=10)
big = Font(name=ARIAL, bold=True, size=12)
brass = Font(name=ARIAL, bold=True, size=10, color="A8853C")

ix["A1"] = "LifeTogether — Grand Master Title Library"; ix["A1"].font = tf
ix["A2"] = "Christian Intelligence(TM) · consolidated, stable-ID'd, deduplicated"
ix["A2"].font = Font(name=ARIAL, italic=True, size=10, color="7C8A80")
ix["A4"] = "Library totals"; ix["A4"].font = h2
r = 5
for name, f in [
    ("Total titles (all collections)", f"=COUNTA({IDR})"),
    ("Unique titles (canonical)", f'=COUNTIF({CANR},"Yes")'),
    ("Duplicate instances (non-canonical)", f'=COUNTIF({CANR},"No")'),
    ("Distinct titles that recur (2+ times)", f'=COUNTIFS({CANR},"Yes",{OCCR},">1")'),
]:
    ix[f"A{r}"] = name; ix[f"A{r}"].font = lbl
    ix[f"C{r}"] = f; ix[f"C{r}"].font = big
    r += 1
r += 1
ix[f"A{r}"] = "By collection"; ix[f"A{r}"].font = h2; r += 1
ix[f"A{r}"], ix[f"B{r}"], ix[f"C{r}"], ix[f"D{r}"] = "Collection", "Code", "Titles", "Unique (canonical)"
for cc in "ABCD": ix[f"{cc}{r}"].font = brass
r += 1
for coll in ORDER:
    ix[f"A{r}"] = coll; ix[f"A{r}"].font = lbl
    ix[f"B{r}"] = CODES[coll]; ix[f"B{r}"].font = lbl
    ix[f"C{r}"] = f"=COUNTIF({COLLR},A{r})"; ix[f"C{r}"].font = lbl
    ix[f"D{r}"] = f'=COUNTIFS({COLLR},A{r},{CANR},"Yes")'; ix[f"D{r}"].font = lbl
    r += 1
ix[f"A{r}"] = "TOTAL"; ix[f"A{r}"].font = big
ix[f"C{r}"] = f"=COUNTA({IDR})"; ix[f"C{r}"].font = big
ix[f"D{r}"] = f'=COUNTIF({CANR},"Yes")'; ix[f"D{r}"].font = big
r += 2
ix[f"A{r}"] = "Most-recurring titles (dedupe hotspots)"; ix[f"A{r}"].font = h2; r += 1
ix[f"A{r}"], ix[f"B{r}"], ix[f"C{r}"] = "Title", "Times", "In # Collections"
for cc in "ABC": ix[f"{cc}{r}"].font = brass
r += 1
for d in top_dups:
    ix[f"A{r}"] = d["title"]; ix[f"A{r}"].font = lbl
    ix[f"B{r}"] = d["occ"]; ix[f"B{r}"].font = lbl
    ix[f"C{r}"] = d["colls"]; ix[f"C{r}"].font = lbl
    r += 1
ix.column_dimensions["A"].width = 40
for cc, w in [("B", 12), ("C", 16), ("D", 18)]:
    ix.column_dimensions[cc].width = w

wb.save(OUT)
print("saved", OUT)
print(f"total={total} unique={unique} dup_instances={dup_instances} recurring={recurring} cross_collection={cross}")
cc = Counter(r["collection"] for r in records)
for coll in ORDER:
    print(f"   {cc.get(coll,0):>4}  {CODES[coll]}  {coll}")
print("top dups:", [(d['title'], d['occ']) for d in top_dups[:8]])
