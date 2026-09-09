"""Re-present the master title library in the format of Five_Biblical_Purposes_Series_Purpose_Driven.xlsx:
one tab per track, columns # / Category / Title / Subtitle / Format, navy header, cream first row,
Arial 9-10, frozen header, and a Summary tab with grouped counts and a grand total."""
import openpyxl, re
from collections import defaultdict, Counter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = "/home/claude/mtl/Master_Title_Library_v20.xlsx"
OUT = "/home/claude/mtl/LifeTogether_Title_Catalog_by_Track.xlsx"

NAVY, CREAM, GOLD, GREY, SAGE = "101E38", "F1EBDD", "B98D3E", "999999", "6E7568"

wb_in = openpyxl.load_workbook(SRC, read_only=True)
ws_in = wb_in["MASTER LIBRARY"]
hdr = [c.value for c in next(ws_in.iter_rows(min_row=1, max_row=1))]
C = {h: i for i, h in enumerate(hdr)}

rows = []
for r in ws_in.iter_rows(min_row=2, values_only=True):
    if not r[0] or r[C["Keep / Review / Archive"]] == "Archive":
        continue
    rows.append(r)
print("rows carried over (Archive excluded):", len(rows))

# ---- track = the topic it teaches, which is what his tabs are keyed on
TRACK_ORDER = [
    ("GROUP 1 — THE FIVE PURPOSES", ["Worship", "Small Groups", "Discipleship", "Church Health", "Spiritual Growth"]),
    ("GROUP 2 — MONEY, STEWARDSHIP & GENEROSITY", ["Generosity", "Stewardship", "Money"]),
    ("GROUP 3 — FAMILY & LEGACY", ["Family Legacy", "Marriage", "Parenting", "Relationships"]),
    ("GROUP 4 — MARKETPLACE & LEADERSHIP", ["Business", "Leadership"]),
    ("GROUP 5 — WHOLE LIFE", ["Purpose", "Health & Flourishing", "Prayer", "Unclear"]),
]
TRACK_TAB = {"Worship": "Worship", "Small Groups": "Fellowship", "Discipleship": "Discipleship",
             "Church Health": "Ministry & Church", "Spiritual Growth": "Mission & Formation",
             "Generosity": "Generosity", "Stewardship": "Stewardship", "Money": "Finances",
             "Family Legacy": "Family Legacy", "Marriage": "Marriage", "Parenting": "Parenting",
             "Relationships": "Relationships", "Business": "Marketplace", "Leadership": "Leadership",
             "Purpose": "Purpose & Calling", "Health & Flourishing": "Flourishing",
             "Prayer": "Prayer", "Unclear": "Unsorted"}

DAYS = re.compile(r"\b(\d{1,3})[-\s]?Days?\b", re.I)
def fmt_of(row):
    t = str(row[3] or "") + " " + str(row[C["Content Type"]] or "") + " " + str(row[C["Level"]] or "")
    m = DAYS.search(t)
    if m: return f"{m.group(1)}-Day"
    ct = str(row[C["Content Type"]] or "")
    return {"Campaign": "Campaign", "Journey": "Journey", "Curriculum": "Curriculum", "Session": "Session",
            "Devotional": "Devotional", "Assessment": "Assessment", "Tool": "Tool", "Training": "Training",
            "Sermon Series": "Sermon Series", "Website / Brand": "Platform", "Business Concept": "Concept",
            "Resource": "Resource", "Book": "Book"}.get(ct, ct or "—")

by_track = defaultdict(list)
for r in rows:
    topic = str(r[C["Topic / Category"]] or "Unclear")
    by_track[TRACK_TAB.get(topic, "Unsorted")].append(r)

wb = openpyxl.Workbook()
summary = wb.active
summary.title = "Summary"

def header_row(ws, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 15
    ws.freeze_panes = "A2"

COLS = ["#", "Category", "Title", "Subtitle", "Format", "ID"]
WIDTHS = [4, 30, 40, 46, 13, 11]
counts = {}
for tab, items in sorted(by_track.items(), key=lambda kv: -len(kv[1])):
    ws = wb.create_sheet(tab[:31])
    header_row(ws, COLS, WIDTHS)
    items.sort(key=lambda r: (str(r[C["Parent / Belongs To"]] or "~"), str(r[3] or "")))
    last_cat = None
    for i, r in enumerate(items, start=1):
        cat = str(r[C["Parent / Belongs To"]] or "").strip()
        shown = cat if cat != last_cat else None          # his convention: category shown once per block
        last_cat = cat
        vals = [i, shown, r[3], r[4], fmt_of(r), r[0]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=i + 1, column=c, value=v)
            cell.font = Font(name="Arial", size=9 if c in (1, 5, 6) else 9.5,
                             color=GREY if c in (1, 6) else None,
                             bold=(c == 2 and shown is not None))
            cell.alignment = Alignment(vertical="center")
            if i == 1:
                cell.fill = PatternFill("solid", fgColor=CREAM)
        ws.row_dimensions[i + 1].height = 15
    counts[tab] = len(items)
    print(f"  {tab:22s} {len(items):6,d}")

# ---------------- Summary, in the same voice as his
summary.column_dimensions["A"].width = 46
summary.column_dimensions["B"].width = 14
summary.column_dimensions["C"].width = 58
summary["A1"] = "LifeTogether Master Title Catalog — by Track"
summary["A1"].font = Font(name="Arial", size=13, bold=True, color=NAVY)
summary["A2"] = ("Every title in the master library that has not been archived, laid out one track per tab. "
                 "Sixteen sources, duplicates collapsed, sessions and days held separately.")
summary["A2"].font = Font(name="Arial", size=10.5, color=SAGE)

r = 4
row_refs = []
for group, topics in TRACK_ORDER:
    cell = summary.cell(row=r, column=1, value=group)
    cell.font = Font(name="Arial", size=10.5, bold=True, color=GOLD)
    r += 1
    for topic in topics:
        tab = TRACK_TAB[topic]
        if tab not in counts: continue
        summary.cell(row=r, column=1, value=tab).font = Font(name="Arial", size=10)
        summary.cell(row=r, column=2, value=counts.pop(tab)).font = Font(name="Arial", size=10)
        note = f"Topic: {topic}" if topic != "Unclear" else "Topic not established by the source"
        summary.cell(row=r, column=3, value=note).font = Font(name="Arial", size=10, color=SAGE)
        row_refs.append(r)
        r += 1
    r += 1
if counts:
    summary.cell(row=r, column=1, value="OTHER TRACKS").font = Font(name="Arial", size=10.5, bold=True, color=GOLD)
    r += 1
    for tab, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        summary.cell(row=r, column=1, value=tab).font = Font(name="Arial", size=10)
        summary.cell(row=r, column=2, value=n).font = Font(name="Arial", size=10)
        row_refs.append(r); r += 1
    r += 1
summary.cell(row=r, column=1, value="GRAND TOTAL").font = Font(name="Arial", size=11, bold=True)
summary.cell(row=r, column=2, value="+".join(f"B{x}" for x in row_refs)).value = "=" + "+".join(f"B{x}" for x in row_refs)
summary.cell(row=r, column=2).font = Font(name="Arial", size=11, bold=True)
r += 2
summary.cell(row=r, column=1, value="How to read this").font = Font(name="Arial", size=10.5, bold=True, color=GOLD)
summary.cell(row=r + 1, column=1,
             value=("Category repeats only at the top of each block, as in your Five Biblical Purposes file. "
                    "Format is taken from the title where it states a length, otherwise from the content type. "
                    "ID matches the master library workbook so any row can be traced back to its source."))
summary.cell(row=r + 1, column=1).font = Font(name="Arial", size=10, color=SAGE)
summary.freeze_panes = "A4"

wb.save(OUT)
print("saved", OUT)
