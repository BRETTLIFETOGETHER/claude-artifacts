import openpyxl, json, re
from collections import Counter

F = "/mnt/user-data/uploads/Five_Biblical_Purposes_Series_Purpose_Driven.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)

PURPOSE = ["Worship", "Fellowship", "Discipleship", "Ministry", "Mission", "Overview"]
THEMED = ["Generosity", "Finances"]                    # theme block, first row is the theme description
SUBBED = ["Stewardship"]                               # theme | title | subtitle | format
CATSUB = ["Family Legacy (UHNW)", "Church Campaigns (General)"]
OCCASION = ["Occasion Finance Messages"]               # title row then subtitle row

rows, notes = [], Counter()


def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


for sheet in PURPOSE:
    ws = wb[sheet]
    for r in range(2, ws.max_row + 1):
        cat, t, days = clean(ws.cell(r, 2).value), clean(ws.cell(r, 3).value), clean(ws.cell(r, 4).value)
        if not t:
            continue
        rows.append(dict(sheet=sheet, cat=cat, title=t, sub="", fmt=(f"{days}-Day" if days else ""), n=r - 1))

for sheet in THEMED:
    ws = wb[sheet]
    theme, theme_desc = "", ""
    for r in range(2, ws.max_row + 1):
        th, t, fmt = clean(ws.cell(r, 2).value), clean(ws.cell(r, 3).value), clean(ws.cell(r, 4).value)
        if th:
            theme = th
        if not fmt:                                     # the block's description line, not a title
            theme_desc = t
            notes["theme description rows skipped"] += 1
            continue
        if t:
            rows.append(dict(sheet=sheet, cat=theme, title=t, sub=theme_desc, fmt=fmt, n=r - 1))

for sheet in SUBBED:
    ws = wb[sheet]
    theme = ""
    for r in range(2, ws.max_row + 1):
        th, t, sub, fmt = (clean(ws.cell(r, c).value) for c in (2, 3, 4, 5))
        if th:
            theme = th
        if t:
            rows.append(dict(sheet=sheet, cat=theme, title=t, sub=sub, fmt=fmt, n=r - 1))

for sheet in CATSUB:
    ws = wb[sheet]
    cat = ""
    for r in range(2, ws.max_row + 1):
        c2, t, sub, fmt = (clean(ws.cell(r, c).value) for c in (2, 3, 4, 5))
        if c2:
            cat = c2
        if t:
            rows.append(dict(sheet=sheet, cat=cat, title=t, sub=sub, fmt=fmt, n=r - 1))

for sheet in OCCASION:
    ws = wb[sheet]
    cat = ""
    r = 2
    while r <= ws.max_row:
        c2, t, fmt = clean(ws.cell(r, 2).value), clean(ws.cell(r, 3).value), clean(ws.cell(r, 4).value)
        if c2:
            cat = c2
        if t and fmt:
            nxt = clean(ws.cell(r + 1, 3).value) if r + 1 <= ws.max_row else ""
            nxt_fmt = clean(ws.cell(r + 1, 4).value) if r + 1 <= ws.max_row else ""
            sub = nxt if (nxt and not nxt_fmt) else ""
            rows.append(dict(sheet=sheet, cat=cat, title=t, sub=sub, fmt=fmt, n=r - 1))
            r += 2 if sub else 1
            continue
        r += 1

# quality problems worth flagging back
double_the = [x for x in rows if re.match(r"^The The\b", x["title"])]
the_prefixed = [x for x in rows if re.match(r"^The [a-z]", x["title"])]
lower_start = [x for x in rows if re.match(r"^\d+ Days? (of|to) [a-z]", x["title"])]
print("rows:", len(rows), "| by sheet:", Counter(x["sheet"] for x in rows).most_common())
print("with subtitle:", sum(1 for x in rows if x["sub"]))
print("skipped:", dict(notes))
print("'The The' artifacts:", len(double_the), "| 'The ' + lowercase:", len(the_prefixed),
      "| 'N Days of ' + lowercase fragment:", len(lower_start))
print("unique titles:", len({x['title'].lower() for x in rows}))
json.dump(rows, open("/home/claude/mtl/rows17.json", "w"))
for x in rows[:3] + [y for y in rows if y["sheet"] == "Stewardship"][:3] + [y for y in rows if y["sheet"] == "Occasion Finance Messages"][:3]:
    print(f"  [{x['sheet'][:18]:20s}] {x['title'][:42]:44s} | {x['sub'][:34]:36s} | {x['fmt']}")
