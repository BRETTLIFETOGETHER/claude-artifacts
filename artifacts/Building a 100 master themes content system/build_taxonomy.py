from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "1B2A4A"; GOLD = "C9A227"; LIGHT = "F5F1E8"; WHITE = "FFFFFF"
HDR_FONT = Font(name="Arial", bold=True, color=WHITE, size=11)
BODY = Font(name="Arial", size=10)
CAT_FILL = PatternFill("solid", start_color=LIGHT)
HDR_FILL = PatternFill("solid", start_color=NAVY)
thin = Side(style="thin", color="D8D2C4")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet_setup(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HDR_FONT; c.fill = HDR_FILL
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

def fill_rows(ws, rows):
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = BODY; cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

wb = Workbook()

# ---------------- THEMES (100) ----------------
themes = [
("Identity & Purpose", [
 ("Made for This","I don't know why I'm here","Ephesians 2:10", "Y"),
 ("Who You Are","I don't know who I am anymore","1 Peter 2:9", "Y"),
 ("Significance","Does my life matter?","Psalm 139:13-14", "Y"),
 ("Purpose","What am I supposed to do with my life?","Jeremiah 29:11", "Y"),
 ("Calling","Is there more than this?","Romans 12:6-8", "Y"),
 ("Enough","I never measure up","2 Corinthians 12:9", "N"),
 ("Known","Nobody really knows me","Psalm 139:1-4", "N"),
 ("Chosen","I feel overlooked","John 15:16", "N"),
 ("Wonderfully Made","I don't like who I see in the mirror","Psalm 139:14", "N"),
 ("Your One Life","I'm afraid of wasting my life","Psalm 90:12", "N")]),
("Peace & Emotional Health", [
 ("Peace","My mind never stops racing","John 14:27", "Y"),
 ("Unshaken","Everything feels unstable","Psalm 16:8", "Y"),
 ("Rest","I'm exhausted all the time","Matthew 11:28-30", "Y"),
 ("Fear Not","Fear is running my decisions","Isaiah 41:10", "Y"),
 ("Anxious for Nothing","Anxiety controls my days","Philippians 4:6-7", "Y"),
 ("Steady","My emotions swing with my circumstances","Psalm 112:7", "N"),
 ("Whole","I feel broken inside","Psalm 147:3", "N"),
 ("Renewed","I need a fresh start in my thinking","Romans 12:2", "N"),
 ("Calm in the Storm","My life is in crisis right now","Mark 4:39", "N"),
 ("Breathe","I can't slow down","Psalm 46:10", "N")]),
("Marriage & Relationships", [
 ("Love That Lasts","We're drifting apart","1 Corinthians 13:4-7", "Y"),
 ("Better Together","We feel like roommates","Ecclesiastes 4:9-12", "Y"),
 ("Covenant","We're not sure we'll make it","Malachi 2:14-15", "Y"),
 ("Stronger","Our marriage needs a reset","Song of Songs 8:6-7", "Y"),
 ("Communication","We can't talk without fighting","James 1:19", "Y"),
 ("Healing After Hurt","Trust has been broken","Colossians 3:13", "N"),
 ("Forgiveness","I can't let go of what they did","Ephesians 4:32", "N"),
 ("Friendship","I'm surrounded but lonely","Proverbs 17:17", "N"),
 ("Us","We've stopped dreaming together","Amos 3:3", "N"),
 ("The Marriage Reset","We want to start over, together","Ruth 1:16-17", "N")]),
("Family & Parenting", [
 ("Family Legacy","What will survive me?","Psalm 78:4-7", "Y"),
 ("Passing Down Faith","Will my kids own their faith?","Deuteronomy 6:6-7", "Y"),
 ("Parenting on Purpose","I'm parenting on autopilot","Proverbs 22:6", "Y"),
 ("Grandparent Legacy","How do I bless my grandchildren?","Psalm 71:18", "Y"),
 ("Family Mission","Why does our family exist?","Joshua 24:15", "Y"),
 ("Around the Table","We're never together anymore","Acts 2:46", "N"),
 ("Home","Our house doesn't feel like a refuge","Psalm 127:1", "N"),
 ("Generations","I want faith to outlive me","Psalm 145:4", "N"),
 ("Launch Ready","My kids are leaving soon","Psalm 127:4", "N"),
 ("Raising Faith","Culture is discipling my kids faster than I am","Ephesians 6:4", "N")]),
("Money & Stewardship", [
 ("God Owns It All","Whose money is it, really?","Psalm 24:1", "Y"),
 ("Money Wise","I need practical wisdom with money","Proverbs 3:9-10", "Y"),
 ("Money Made Simple","Money feels overwhelming","Proverbs 21:5", "Y"),
 ("Contentment","I always want more","Philippians 4:11-13", "Y"),
 ("First Things First","My priorities are out of order","Matthew 6:33", "Y"),
 ("Stewardship","I want to manage what God gave me well","Matthew 25:21", "N"),
 ("Provision","I'm afraid there won't be enough","Matthew 6:31-33", "N"),
 ("Open Hands","I hold everything too tightly","1 Timothy 6:17-19", "N"),
 ("More Than Enough","I live in scarcity even in abundance","2 Corinthians 9:8", "N"),
 ("Free Indeed","Debt and money stress own me","Proverbs 22:7", "N")]),
("Generosity & Kingdom Impact", [
 ("Generous Life","I want giving to define me","2 Corinthians 9:6-7", "Y"),
 ("The Joy of Giving","Giving feels like obligation, not joy","Acts 20:35", "Y"),
 ("Kingdom Impact","I want my resources to count forever","Matthew 6:19-21", "Y"),
 ("Give First","Giving is my leftover, not my first fruit","Proverbs 3:9", "Y"),
 ("Overflow","I want to be a channel, not a reservoir","Malachi 3:10", "Y"),
 ("Blessed to Bless","Why has God given us so much?","Genesis 12:2", "N"),
 ("The Giving Family","I want my kids to love giving","Deuteronomy 15:10", "N"),
 ("Legacy of Generosity","I want generosity in my family's DNA","Proverbs 11:24-25", "N"),
 ("Live to Give","I want a lifestyle of generosity","Luke 6:38", "N"),
 ("Irrational Generosity","I want to give in a way that requires faith","Mark 12:41-44", "N")]),
("Faith & Discipleship", [
 ("Following Jesus","What does it actually mean to follow him?","Luke 9:23", "Y"),
 ("Rooted","My faith is a mile wide and an inch deep","Colossians 2:6-7", "Y"),
 ("Abide","I'm working for God but not with God","John 15:4-5", "Y"),
 ("Apprentice","I want to learn to live like Jesus","Matthew 11:29", "Y"),
 ("The Practices","I need habits that form my soul","1 Timothy 4:7-8", "Y"),
 ("The Word","I don't know my Bible","Psalm 119:105", "N"),
 ("Prayer","My prayer life is flat","Luke 11:1", "N"),
 ("Walk This Way","I believe, but my life doesn't show it","Micah 6:8", "N"),
 ("Grow","I've been stuck spiritually for years","2 Peter 3:18", "N"),
 ("Fasting and Feasting","I want to hunger for God again","Matthew 5:6", "N")]),
("Community & Belonging", [
 ("Never Alone","I'm doing life by myself","Ecclesiastes 4:9-10", "Y"),
 ("Belong","I don't fit anywhere","Romans 12:5", "Y"),
 ("Group Up","I want people to do life with","Hebrews 10:24-25", "Y"),
 ("One Another","I attend church but I'm not connected","John 13:34-35", "Y"),
 ("Life Together","I want real community, not small talk","Acts 2:42-47", "Y"),
 ("Connected","I'm surrounded by people and still lonely","1 Corinthians 12:26-27", "N"),
 ("Welcome Home","I want our church to feel like family","Romans 15:7", "N"),
 ("Doing Life Together","I want friendships that go deep","1 Thessalonians 2:8", "N"),
 ("The Church We Long For","I've been hurt by church","Ephesians 4:2-3", "N"),
 ("Come to the Table","I want hospitality to mark our home","Luke 14:13-14", "N")]),
("Hope & Trials", [
 ("Hope Rising","I've lost hope","Romans 15:13", "Y"),
 ("Unbreakable","Life keeps knocking me down","2 Corinthians 4:8-9", "Y"),
 ("When Life Hurts","Why is this happening to me?","Romans 8:28", "Y"),
 ("Grief to Grace","I'm walking through loss","Psalm 34:18", "Y"),
 ("Trust in the Dark","I can't see what God is doing","Proverbs 3:5-6", "Y"),
 ("Comeback","I need to believe again after failure","Joel 2:25", "N"),
 ("Redeemed","My past disqualifies me","2 Corinthians 5:17", "N"),
 ("New Every Morning","I need daily mercy, not one-time answers","Lamentations 3:22-23", "N"),
 ("Anchor","I need something that holds","Hebrews 6:19", "N"),
 ("Light in the Darkness","This season feels endless","Psalm 30:5", "N")]),
("Mission & Legacy", [
 ("Sent","My faith is for me but not through me","John 20:21", "Y"),
 ("Love Your City","I want our church to matter to our town","Jeremiah 29:7", "Y"),
 ("Every Neighbor","I don't know the people next door","Luke 10:27", "Y"),
 ("What Matters Most","I'm busy with things that won't last","Matthew 16:26", "Y"),
 ("Finish Strong","The last chapters should be the best","2 Timothy 4:7", "Y"),
 ("Go","I've never shared my faith","Matthew 28:19-20", "N"),
 ("The Blessing","I want to speak life over the next generation","Numbers 6:24-26", "N"),
 ("On Mission","I want my everyday life to count","Colossians 3:23-24", "N"),
 ("Lasting Legacy","I want to leave more than money","Proverbs 13:22", "N"),
 ("Commissioned","It's time to send the next generation","Joshua 1:9", "N")]),
]

ws = wb.active; ws.title = "Themes"
sheet_setup(ws, ["ID","Theme","Category","Felt Need","Scripture Anchor","Core 50"], [8,26,26,44,22,10])
rows = []; n = 1
for cat, items in themes:
    for name, need, verse, core in items:
        rows.append((f"T{n:03d}", name, cat, need, verse, core)); n += 1
fill_rows(ws, rows)

# ---------------- AUDIENCES (100) ----------------
aud = [
("Life Stage", ["Children","Preteens","Middle school students","High school students","College students","Young adults","Singles","Adults in their 30s-40s","Midlife adults","Empty nesters","Pre-retirees","Retirees","Seniors","Widows and widowers","Young professionals"]),
("Family & Marriage", ["Engaged couples","Newlyweds","Married couples","Couples in crisis","Parents of infants","Parents of toddlers","Parents of school-age kids","Parents of teens","Parents of adult children","Single parents","Blended families","Adoptive families","Foster families","Grandparents","Multi-generational families"]),
("Church Life", ["Seekers","New believers","New members","Sunday-only attenders","Small group members","Small group leaders","Group hosts","Volunteers","Ministry team leaders","Worship teams","Church staff","Elders and deacons","Senior pastors","Church planters","Missionaries"]),
("Marketplace", ["Business owners","Entrepreneurs","Executives","Managers","Employees","Remote workers","Healthcare workers","Teachers and educators","First responders","Military families","Veterans","Farmers and ranchers","Tradespeople","Creatives and artists","Athletes and coaches"]),
("Situational", ["Those battling anxiety","Those in a season of depression","The grieving","Those in recovery","The unemployed","Those in career transition","The recently relocated","Caregivers","Those facing illness","The lonely","The doubting and deconstructing","Returning prodigals","The divorced","Families of the incarcerated","The financially stressed","New homeowners","New graduates","Those facing aging parents","Those entering retirement transition","Those with sudden new wealth"]),
("Legacy & Wealth", ["High-net-worth families","Ultra-high-net-worth families","Legacy families","Family business owners","Family offices","Heirs and the rising generation","Business sellers","Kingdom investors","Major donors","Family foundations","Giving circles","Donor-advised fund holders","Estate planning families","Multi-generational wealth families","Christian advisors' client families","Widowed wealth stewards","Inheritors","Philanthropic couples","Succession-stage owners","Retiring founders"]),
]
ws = wb.create_sheet("Audiences")
sheet_setup(ws, ["ID","Audience Segment","Category"], [8,42,24])
rows = []; n = 1
for cat, items in aud:
    for a in items:
        rows.append((f"A{n:03d}", a, cat)); n += 1
fill_rows(ws, rows)

# ---------------- LIFE EVENTS (100, Brett's list verbatim) ----------------
events = [
("Family Beginnings", ["Engagement","Wedding","Honeymoon / First Year of Marriage","Purchasing First Home","Birth of First Child","Adoption","Foster Parenting","Blended Family Formation","Child Dedication","Baptism of a Child"]),
("Raising Children", ["Child Starts Kindergarten","First Allowance","First Job","First Bank Account","First Investment Account","Child Gets Driver's License","First Mission Trip","First Summer Job","High School Graduation","College Decision"]),
("Young Adult Years", ["Leaving Home","College Graduation","First Career","Graduate School","Engagement of Adult Child","Wedding of Adult Child","Birth of First Grandchild","Launching a Business","Purchasing First Home (Adult Child)","Becoming Financially Independent"]),
("Family Milestones", ["Anniversary Milestones","Becoming Grandparents","Family Reunion","Multi-Generational Vacation","Family Retreat","Writing Family History","Creating Family Mission Statement","Creating Family Constitution","First Family Council","Annual Family Meeting"]),
("Financial Events", ["Significant Salary Increase","Bonus or Windfall","Inheritance Received","Selling a Business","Buying a Business","IPO or Liquidity Event","Major Investment Gain","Significant Financial Loss","Paying Off Debt","Reaching Financial Independence"]),
("Business Events", ["Starting a Business","Hiring First Employee","Leadership Transition","Adding Family Members to Business","Succession Planning","Business Sale","Retirement from Business","Merging Companies","Business Crisis","Business Recovery"]),
("Generosity & Kingdom Impact", ["First Major Gift","Creating a Giving Plan","Opening a Donor-Advised Fund","Starting a Family Foundation","Joining a Giving Circle","Funding a Ministry Project","Mission Trip as a Family","Estate Gift Planning","Legacy Giving Decision","Family Service Project"]),
("Estate & Legacy", ["Writing a Will","Creating a Trust","Updating Estate Plan","Selecting Trustees","Naming Guardians","Family Estate Meeting","Writing Ethical Will","Writing Legacy Letters","Family Wealth Conversation","Preparing Responsible Heirs"]),
("Aging & Retirement", ["Retirement Planning","Retirement","Downsizing Home","Becoming a Caregiver","Parent Moves Into Assisted Living","Long-Term Care Decision","Serious Health Diagnosis","Hospice Care","Death of a Parent","Death of a Spouse"]),
("Spiritual Milestones", ["Coming to Faith in Christ","Baptism","Joining a Church","Becoming a Small Group Leader","Beginning Family Devotions","Launching a Family Ministry","Sabbatical or Spiritual Retreat","Commissioning the Next Generation","Celebrating a Family Legacy","End-of-Life Blessing & Celebration"]),
]
ws = wb.create_sheet("LifeEvents")
sheet_setup(ws, ["ID","Life Event","Category"], [8,44,28])
rows = []; n = 1
for cat, items in events:
    for e in items:
        rows.append((f"E{n:03d}", e, cat)); n += 1
fill_rows(ws, rows)

# ---------------- TRUSTED GUIDES (50) ----------------
guides = [
("Church", ["Senior pastor","Executive pastor","Small groups pastor","Discipleship pastor","Youth pastor","Children's pastor","Women's ministry leader","Men's ministry leader","Missions pastor","Worship pastor","Church planter","Denominational leader"]),
("Financial & Legal", ["Financial advisor","Wealth manager","Family office advisor","Estate attorney","CPA","Insurance advisor","Planned giving officer","Foundation officer","Donor-advised fund sponsor","Kingdom investing advisor","Business exit advisor","Philanthropy consultant"]),
("Family & Coaching", ["Christian counselor","Marriage and family therapist","Life coach","Business coach","Legacy coach","Mentor","Grandparent","Parent","Small group leader","Family meeting facilitator"]),
("Marketplace & Institutions", ["Workplace ministry leader","Corporate chaplain","Military or hospital chaplain","CEO or marketplace leader","Christian school leader","Homeschool co-op leader","Camp or retreat director","Nonprofit ministry director"]),
("Media & Network", ["Author","Publisher","Podcaster","Radio host","Conference host","Ministry influencer","Parachurch leader","Giving circle leader"]),
]
ws = wb.create_sheet("TrustedGuides")
sheet_setup(ws, ["ID","Trusted Guide","Category"], [8,38,28])
rows = []; n = 1
for cat, items in guides:
    for g in items:
        rows.append((f"G{n:03d}", g, cat)); n += 1
fill_rows(ws, rows)

# ---------------- FORMATS (25) ----------------
formats = ["Seven Day Experience","21-Day Reset","30-Day Devotional","40 Day Campaign","6-Session Small Group Study","4-6 Week Sermon Series","8-Week Course","Weekend Retreat","One-Day Workshop","Family Meeting Toolkit","Assessment + Report","Legacy Letter Builder","Video Series","Audio Devotional / Podcast","Email Devotional","Text (SMS) Journey","Printed Devotional Book","Guided Journal","Participant Workbook","Facilitator Guide","Children's Edition","Student Edition","Couples Edition","Advisor-Led Cohort","Custom Family Journey"]
ws = wb.create_sheet("Formats")
sheet_setup(ws, ["ID","Format"], [8,40])
fill_rows(ws, [(f"F{i:03d}", f) for i, f in enumerate(formats, 1)])

# ---------------- FRAMEWORKS (20) ----------------
frameworks = [
("The 40-Day Campaign Arc","Kickoff Sunday through Celebration Sunday with daily Scripture, devotional, step, and spiritual partner conversation"),
("Seven Day Experience Arc","Kickoff, Days 1-6, Celebration Sunday"),
("Story Engine","Testimony gathering pipeline that makes Celebration Sunday the congregation showing each other what happened"),
("Celebration Sunday Model","The ministry invitation as the natural, fully prepared conclusion of the campaign"),
("Spiritual Partner Model","Every participant paired for daily conversation and accountability"),
("TNT Teleprompter Format","Pastor-led daily video scripts"),
("Group Up Seasonal Community","People group up for a season and become relational discipleship communities"),
("HOST Group Launch Model","Anyone with a Home, Outreach heart, Snacks, and a TV can host a group"),
("Felt-Need Title Logic","The title names the felt need in language the congregation already uses"),
("Four-Part Blueprint","Campaign Snapshot, 6-Session Curriculum, 30/40-Day Journey, sample devotional and session"),
("Six-Index Taxonomy Tagging","Every resource tagged by Affinity, Guide, Biblical Topic, Felt Need, Life Question, Life Event"),
("Life-Event Trigger Recommendations","The platform recommends pathways based on what just happened in a family's life"),
("Campaign Builder Questionnaire","6-10 questions that generate a recommended title, outline, and pathway"),
("Family Meeting Framework","Agenda, questions, and tools for a facilitated family conversation"),
("Legacy Letter Framework","Guided writing of what your family needs to hear from you"),
("Heir Readiness Pathway","Assessment and staged preparation of the rising generation"),
("Giving Plan Framework","From first major gift to multi-generational giving strategy"),
("Custom Family Journey","8 steps from assessment to annual family stewardship review"),
("Advisor Co-Brand Gifting Flow","Advisor selects, co-brands, and gifts a resource to a client family in minutes"),
("Annual Family Stewardship Review","Yearly rhythm that keeps the family pathway alive"),
]
ws = wb.create_sheet("Frameworks")
sheet_setup(ws, ["ID","Framework","Description"], [8,34,80])
fill_rows(ws, [(f"K{i:03d}", n_, d) for i, (n_, d) in enumerate(frameworks, 1)])

# ---------------- CAMPAIGNS (schema + Theme seeded with 50) ----------------
fl = [
("40 Days of Family Legacy","What Will Survive You","High-net-worth families","Family Estate Meeting"),
("What Matters Most","A 40-Day Journey to the Life Behind the Wealth","Midlife adults","Updating Estate Plan"),
("Lasting Wealth","Passing Down More Than Money","Legacy families","Family Wealth Conversation"),
("Preparing the Ones You Love","Getting Heirs Ready for What Is Coming","Parents of adult children","Preparing Responsible Heirs"),
("The Family Table","40 Days of Conversations That Change Generations","Multi-generational families","Annual Family Meeting"),
("Generations","Faith That Outlives You","Grandparents","Becoming Grandparents"),
("The Blessing","Speaking Life Over Your Children and Grandchildren","Grandparents","Birth of First Grandchild"),
("Grandparent Legacy","40 Days to Bless the Next Generation","Grandparents","Birth of First Grandchild"),
("Passing Down Faith","The Inheritance Money Cannot Buy","Parents of teens","Beginning Family Devotions"),
("Our Family Story","Discovering Where You Came From and Why It Matters","Multi-generational families","Writing Family History"),
("The Family Mission","40 Days to Discover Why Your Family Exists","Married couples","Creating Family Mission Statement"),
("House of Values","Naming What Your Family Stands For","Legacy families","Creating Family Constitution"),
("The First Family Meeting","Starting the Conversation You Have Been Avoiding","Family business owners","First Family Council"),
("The Responsible Heir","Raising Children Who Can Handle What Is Coming","Heirs and the rising generation","Preparing Responsible Heirs"),
("Beyond the Estate Plan","What the Documents Cannot Do","Estate planning families","Writing a Will"),
("Legacy Letters","Writing What Your Family Needs to Hear","Seniors","Writing Legacy Letters"),
("The Ethical Will","Leaving Wisdom, Not Just Wealth","Pre-retirees","Writing Ethical Will"),
("Rooted Together","Building a Family That Stays Close","Parents of school-age kids","Family Retreat"),
("Stronger Roots","40 Days of Family Spiritual Habits","Parents of school-age kids","Beginning Family Devotions"),
("The Long View","Thinking in Generations, Not Quarters","Kingdom investors","Legacy Giving Decision"),
("Handing It Down","The Art of Transferring What Matters","Legacy families","Inheritance Received"),
("Faith of Our Fathers","Honoring the Generation Before You","Adults in their 30s-40s","Death of a Parent"),
("The Family Constitution","Agreements That Hold a Family Together","Family offices","Creating Family Constitution"),
("Unity","Keeping a Family Whole When Wealth Could Divide It","Multi-generational wealth families","Inheritance Received"),
("Peace in the Family","Healing Old Wounds Before It Is Too Late","Multi-generational families","Family Reunion"),
("Reconciled","40 Days Toward Family Forgiveness","Blended families","Blended Family Formation"),
("The Family Reunion","Turning Gatherings Into Legacy Moments","Multi-generational families","Family Reunion"),
("Around One Table","Multi-Generational Faith in Practice","Multi-generational families","Multi-Generational Vacation"),
("Launch Ready","Preparing Your Children to Leave Well","Parents of teens","High School Graduation"),
("Welcome to the Family","Bringing In-Laws Into the Legacy","Parents of adult children","Wedding of Adult Child"),
("The Blended Blessing","Building Legacy in a Blended Family","Blended families","Blended Family Formation"),
("Adopted","The Family God Builds","Adoptive families","Adoption"),
("The Family Business","Faith, Work, and the Name on the Door","Family business owners","Adding Family Members to Business"),
("Succession","Handing Over the Keys Without Losing the Family","Succession-stage owners","Succession Planning"),
("After the Sale","Who We Are When the Business Is Gone","Business sellers","Selling a Business"),
("The Steward's Family","God Owns It All, Together","Legacy families","Family Wealth Conversation"),
("Generous Together","A Family Giving Journey","Philanthropic couples","Creating a Giving Plan"),
("The Giving Family","Raising Children Who Love to Give","Parents of school-age kids","First Allowance"),
("Kingdom Impact","Deploying Your Family for What Lasts","Kingdom investors","Funding a Ministry Project"),
("The Family Foundation","Giving With Purpose Across Generations","Family foundations","Starting a Family Foundation"),
("Grandchildren of Promise","Praying the Next Generation Home","Grandparents","Becoming Grandparents"),
("Empty Nest, Full Purpose","Legacy After the Kids Leave","Empty nesters","Leaving Home"),
("Finishing Strong","The Last Chapters Are the Best Chapters","Retirees","Retirement"),
("Before You Go","Conversations to Have While There Is Time","Caregivers","Parent Moves Into Assisted Living"),
("The Caregiver's Calling","Honoring Parents in Their Final Season","Caregivers","Becoming a Caregiver"),
("Saying Goodbye Well","Grief, Gratitude, and the Hope of Heaven","The grieving","Death of a Parent"),
("The Widow's Strength","Rebuilding Legacy After Loss","Widows and widowers","Death of a Spouse"),
("Heirlooms","The Stories Behind the Things We Keep","Seniors","Downsizing Home"),
("The Family Historian","Capturing a Century Before It Is Lost","Multi-generational families","Writing Family History"),
("Commissioned","Sending the Next Generation Into Their Calling","Parents of adult children","Commissioning the Next Generation"),
]
ws = wb.create_sheet("Campaigns")
headers = ["CampaignID","Theme","Campaign Title","Subtitle","Audience","Life Event","Trusted Guide","Format","Framework","Felt Need","Scripture Anchor","Tier","Status","Doc URL"]
sheet_setup(ws, headers, [12,18,30,44,28,28,22,20,26,36,20,10,12,40])
rows = []
for i, (t, s, a, e) in enumerate(fl, 1):
    rows.append((f"C{i:04d}","Family Legacy",t,s,a,e,"Wealth advisor / Family office / Pastor","40 Day Campaign","The 40-Day Campaign Arc","What will survive me?","Psalm 78:4-7","Flagship" if i<=10 else "Core","Ready",""))
fill_rows(ws, rows)

# ---------------- README ----------------
ws = wb.create_sheet("README", 0)
ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 110
title = ws.cell(row=2, column=2, value="LIFETOGETHER MASTER TAXONOMY — THE PLATFORM DATABASE")
title.font = Font(name="Arial", bold=True, size=16, color=NAVY)
lines = [
"",
"This workbook is the single source of truth for the entire platform. Every campaign, devotional, doc, and web page is generated FROM this data — never the other way around.",
"",
"THE SIX MASTER INDEXES",
"Themes (100) — What biblical truth does it teach? The 'Core 50' column marks the 50 themes for the first doc build (50 themes x 50 campaigns = 2,500 Google Docs).",
"Audiences (100) — Who is this for?",
"LifeEvents (100) — What just happened? The single most important index. People do not wake up wanting a stewardship curriculum; they wake up because life happened.",
"TrustedGuides (50) — Who is introducing it?",
"Formats (25) — What shape does it take?",
"Frameworks (20) — What proven engine drives it?",
"",
"HOW CAMPAIGNS ARE GENERATED",
"A campaign is a combination: Theme x Audience x Life Event x Guide x Format x Framework. 100 x 100 x 100 x 50 x 25 x 20 is billions of combinations — the Campaigns tab is where the curated ones live. Every row becomes one Google Doc (via the DocGenerator script) and one web page (via the site generator).",
"",
"THE DOC PIPELINE",
"1. Fill Campaigns rows (Status = Ready).  2. Upload this file to Google Sheets.  3. Install DocGenerator.gs via Extensions > Apps Script.  4. Run generateDocs — it creates one Google Doc per row from your template, files it in a folder per theme, and writes the Doc URL back into the row. No Keyboard Maestro required.",
"",
"COUNTS",
]
r = 3
for ln in lines:
    c = ws.cell(row=r, column=2, value=ln)
    c.font = Font(name="Arial", size=11, bold=ln.isupper() and len(ln) > 0)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if ln and not ln.isupper(): ws.row_dimensions[r].height = 30
    r += 1
counts = [("Themes","=COUNTA(Themes!A2:A1000)"),("Audiences","=COUNTA(Audiences!A2:A1000)"),("Life Events","=COUNTA(LifeEvents!A2:A1000)"),("Trusted Guides","=COUNTA(TrustedGuides!A2:A1000)"),("Formats","=COUNTA(Formats!A2:A1000)"),("Frameworks","=COUNTA(Frameworks!A2:A1000)"),("Campaigns seeded","=COUNTA(Campaigns!A2:A10000)")]
for name, f in counts:
    ws.cell(row=r, column=2, value=name).font = Font(name="Arial", size=11)
    c = ws.cell(row=r, column=3, value=f); c.font = Font(name="Arial", size=11, bold=True, color="0000FF")
    r += 1
ws.column_dimensions["C"].width = 12

wb.save("/home/claude/taxonomy/Lifetogether_Master_Taxonomy.xlsx")
print("saved")
