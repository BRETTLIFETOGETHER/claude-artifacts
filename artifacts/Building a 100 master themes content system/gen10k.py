# -*- coding: utf-8 -*-
import json, sys
sys.path.insert(0,'/home/claude/site')
from data_a import THEMES_A
from data_b import THEMES_B
from data_new import THEMES_NEW, EDW_OLD, AUDS35, EVTS30, SEASON5, FILL25

CATS={1:'Identity & Purpose',2:'Peace & Emotional Health',3:'Marriage & Relationships',4:'Family & Parenting',
5:'Money & Stewardship',6:'Generosity & Kingdom Impact',7:'Faith & Discipleship',8:'Community & Belonging',
9:'Hope & Trials',10:'Mission & Legacy'}

# unify: (theme,cat,scr,need,W,cores[])
ALL=[]
for th,cat,scr,need,cores in (THEMES_A+THEMES_B):
    ALL.append((th,cat,scr,need,EDW_OLD[th],cores))
for th,cat,scr,need,W,cores in THEMES_NEW:
    ALL.append((th,cat,scr,need,W,cores))
assert len(ALL)==100, len(ALL)
Ws=[t[4] for t in ALL]; assert len(set(Ws))==100, 'edition words not unique'
# order themes grouped by category (old 5 then new 5 per category)
ALL.sort(key=lambda t:(t[1], 0 if t[0] in EDW_OLD else 1))

AUD_ROT={1:['Young adults','College students','Midlife adults','New graduates','Those in career transition','Adults in their 30s-40s'],
2:['Those battling anxiety','Caregivers','Parents of teens','Church staff','Healthcare workers','The weary'],
3:['Engaged couples','Newlyweds','Married couples','Couples in crisis','Empty nesters','Blended families'],
4:['Parents of school-age kids','Parents of teens','Grandparents','Multi-generational families','Parents of adult children','Legacy families'],
5:['Young professionals','Married couples','Business owners','The financially stressed','Pre-retirees','New homeowners'],
6:['Major donors','Kingdom investors','Philanthropic couples','Family foundations','Giving circles','Generous families'],
7:['New believers','Small group members','Small group leaders','Sunday-only attenders','Volunteers','Church staff'],
8:['The lonely','New members','The recently relocated','Singles','Young adults','Seniors'],
9:['The grieving','Those facing illness','Those in recovery','The unemployed','Widows and widowers','Caregivers'],
10:['Retirees','Empty nesters','Legacy families','Everyday believers','Business owners','Families on mission']}
EVT_ROT={1:['College Graduation','First Career','Leaving Home','Coming to Faith in Christ','Baptism','Retirement'],
2:['Serious Health Diagnosis','Becoming a Caregiver','Business Crisis','Significant Financial Loss','A New Season','Burnout Recovery'],
3:['Engagement','Wedding','First Year of Marriage','Anniversary Milestones','Blended Family Formation','Purchasing First Home'],
4:['Birth of First Child','Child Dedication','High School Graduation','Becoming Grandparents','Annual Family Meeting','Family Mission Statement'],
5:['First Bank Account','Purchasing First Home','Paying Off Debt','Salary Increase','Retirement Planning','Bonus or Windfall'],
6:['First Major Gift','Creating a Giving Plan','Opening a Donor-Advised Fund','Starting a Family Foundation','Legacy Giving Decision','Funding a Ministry Project'],
7:['Coming to Faith in Christ','Baptism','Joining a Church','Becoming a Group Leader','Beginning Family Devotions','Spiritual Retreat'],
8:['Joining a Church','Becoming a Group Leader','Family Reunion','Relocation','Retirement','Launching a Ministry'],
9:['Death of a Parent','Death of a Spouse','Serious Health Diagnosis','Significant Financial Loss','Business Crisis','A Long Recovery'],
10:['Retirement','Commissioning the Next Generation','Writing Legacy Letters','Succession Planning','Celebrating a Family Legacy','A New Mission Field']}
GUIDE={1:['Senior pastor','Life coach','Mentor'],2:['Christian counselor','Senior pastor','Small group leader'],
3:['Marriage and family therapist','Senior pastor','Christian counselor'],4:['Parent','Grandparent','Family meeting facilitator'],
5:['Financial advisor','CPA','Senior pastor'],6:['Wealth manager','Planned giving officer','Foundation officer'],
7:['Discipleship pastor','Small group leader','Mentor'],8:['Small groups pastor','Small group leader','Senior pastor'],
9:['Christian counselor','Senior pastor','Chaplain'],10:['Missions pastor','Legacy coach','Senior pastor']}
SEAS={1:[['fall'],['any']],2:[['new'],['any']],3:[['any'],['new']],4:[['fall'],['any']],5:[['new'],['yearend']],
6:[['yearend'],['any']],7:[['fall'],['lent']],8:[['fall'],['any']],9:[['lent'],['any']],10:[['any'],['yearend']]}
OUT={1:['discipleship'],2:['emotional'],3:['discipleship'],4:['discipleship'],5:['generosity'],
6:['generosity'],7:['discipleship'],8:['groups'],9:['emotional'],10:['outreach']}
SEASKEY={'for the New Year':['new'],'Through Lent':['lent'],'This Fall':['fall'],'in Advent':['advent'],"at Year's End":['yearend']}
FMT_ROT=['40-Day','21-Day','40-Day','30-Day','40-Day','6-Week','40-Day','7-Day','30-Day','40-Day','21-Day','40-Day','6-Week','30-Day','40-Day','21-Day','40-Day','30-Day','21-Day','40-Day']
DAYS={'7-Day':'seven','21-Day':'21','30-Day':'30','40-Day':'40','6-Week':'six-week'}

SUB_A=["A guided journey in {w} for {aud}","Daily Scripture & one next step for {aud}",
"Where {w} meets the real world of {aud}","{w}, written for {aud}","Walking with {aud} into {w}"]
SUB_E=["A companion for the road {evt}","Daily hope and next steps {evt}","God's Word for life {evt}","{w} for the season you're in"]
SUB_S=["A seasonal journey in {w}","Launch-ready {w} for the church calendar","{w} when the whole church is watching the calendar"]
SUB_F=["A fresh on-ramp into {w}","Practical daily steps in {w}","{w}, one day at a time","A proven path into {w}"]

def stripart(W):
    for a in ('A ','The '):
        if W.startswith(a): return W[len(a):]
    return W

rows=[]; wb_rows=[]; titles=set(); collide=0
cid=0
for th,cat,scr,need,W,cores in ALL:
    camps=[]  # (title,sub,aud_slot,evt_slot,kind)
    for t,s in cores:
        camps.append((t,s,None,None,'core'))
    for i,aud in enumerate(AUDS35):
        t=f"{W} for {aud}"
        camps.append((t,SUB_A[i%len(SUB_A)].format(w=W,aud=aud),aud,None,'aud'))
    for i,evt in enumerate(EVTS30):
        t=f"{W} {evt}"
        low=evt[0].lower()+evt[1:]
        camps.append((t,SUB_E[i%len(SUB_E)].format(w=W,evt=low),None,evt,'evt'))
    for i,sn in enumerate(SEASON5):
        camps.append((f"{W} {sn}",SUB_S[i%len(SUB_S)].format(w=W),None,None,('seas',sn)))
    fi=0
    for pat in FILL25:
        if len(camps)>=100: break
        Wp = stripart(W) if pat.startswith(('7 ','21 ','30 ','40 ')) else W
        camps.append((pat.format(W=Wp),SUB_F[fi%len(SUB_F)].format(w=W),None,None,'fill')); fi+=1
    camps=camps[:100]
    # global dedupe within composition
    final=[]
    for t,s,a,e,k in camps:
        if t in titles:
            for alt in (f"{t} (Group Edition)",f"{t}, Together",f"{t} — A Journey",):
                if alt not in titles: t=alt; collide+=1; break
        titles.add(t); final.append((t,s,a,e,k))
    assert len(final)==100,(th,len(final))
    for i,(t,s,a,e,k) in enumerate(final):
        cid+=1; id_=f"C{cid:05d}"
        aud=a or AUD_ROT[cat][i%6]; evt=e or EVT_ROT[cat][i%6]
        fmt=FMT_ROT[i%20]
        seas=k[1:] and SEASKEY[k[1]] if isinstance(k,tuple) else SEAS[cat][i%2]
        if isinstance(k,tuple): seas=SEASKEY[k[1]]
        out=OUT[cat]
        pop=99-(i*2)%70-(cid%9); nw=(cid*41)%100
        flag=1 if (k=='core' and i==0) or t.startswith('40 Days of') else 0
        rows.append([t,s,cat,fmt,seas,out,pop,nw,flag,id_,th,aud,evt,scr])
        wb_rows.append((id_,th,t,s,aud,evt,GUIDE[cat][i%3],fmt+(' Campaign' if fmt=='40-Day' else ''),
                        'The 40-Day Campaign Arc',need,scr,'Flagship' if flag else 'Core','Ready',''))

print('campaigns:',len(rows),'| title collisions auto-resolved:',collide,'| unique titles:',len(titles)==len(rows))
goia=[r for r in rows if r[0]=='God Owns It All'][0]; print('GOIA id:',goia[9])

META={t[0]:[t[1],t[2],t[3],t[4]] for t in ALL}
with open('/home/claude/site/out/data.js','w') as f:
    f.write('window.CATS='+json.dumps(CATS)+';\n')
    f.write('window.THEMEMETA='+json.dumps(META,ensure_ascii=False)+';\n')
    f.write('window.CAMPAIGNS='+json.dumps(rows,ensure_ascii=False)+';\n')
import os; print('data.js:',os.path.getsize('/home/claude/site/out/data.js'),'bytes')

if __name__=='__main__' and '--wb' in sys.argv:
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    p='/mnt/user-data/outputs/Lifetogether_Master_Taxonomy.xlsx'
    wb=openpyxl.load_workbook(p)
    if 'Campaigns' in wb.sheetnames: del wb['Campaigns']
    ws=wb.create_sheet('Campaigns')
    hdr=['Campaign ID','Theme','Campaign Title','Subtitle','Audience','Life Event','Trusted Guide','Format','Framework','Felt Need','Scripture Anchor','Tier','Status','Doc URL']
    navy=PatternFill('solid',fgColor='1B2A4A'); thin=Border(*[Side(style='thin',color='D9D9D9')]*4)
    for j,h in enumerate(hdr,1):
        c=ws.cell(1,j,h);c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10);c.fill=navy;c.border=thin;c.alignment=Alignment(vertical='center')
    for i,r in enumerate(wb_rows,2):
        for j,v in enumerate(r,1):
            c=ws.cell(i,j,v);c.font=Font(name='Arial',size=10);c.border=thin
    for j,w in enumerate([12,24,34,44,26,30,24,18,24,34,20,10,10,14],1):ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:N{len(wb_rows)+1}'
    wb.save(p);print('workbook: 10,000 rows written')
